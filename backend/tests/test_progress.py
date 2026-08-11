from datetime import datetime

from tests.helpers import auth_headers


def test_progress_me_aggregates_completed_scores(client, student):
    auth = auth_headers(student["token"])
    client.post("/api/scores", json={"game_id": 1, "points": 500, "completed": True}, headers=auth)
    client.post("/api/scores", json={"game_id": 2, "points": 700, "completed": True}, headers=auth)
    client.post("/api/scores", json={"game_id": 1, "points": 999, "completed": False}, headers=auth)

    resp = client.get("/api/progress/me", headers=auth)
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["total_attempts"] == 3
    assert data["total_completed"] == 2
    assert data["total_points"] == 1200
    assert data["distinct_games_completed"] == 2


def test_progress_export_returns_xlsx(client, student):
    auth = auth_headers(student["token"])
    client.post("/api/scores", json={"game_id": 1, "points": 500, "completed": True}, headers=auth)

    resp = client.get("/api/progress/export", headers=auth)
    assert resp.status_code == 200
    assert resp.headers["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert resp.data[:2] == b"PK"  # xlsx dosyaları zip (PK) imzasıyla başlar


def test_progress_me_filters_by_date_range(client, student, app):
    from app.extensions import db
    from app.models import Score

    auth = auth_headers(student["token"])
    client.post("/api/scores", json={"game_id": 1, "points": 500, "completed": True}, headers=auth)
    client.post("/api/scores", json={"game_id": 2, "points": 700, "completed": True}, headers=auth)

    with app.app_context():
        scores = Score.query.order_by(Score.id).all()
        scores[0].created_at = datetime(2024, 1, 15)
        scores[1].created_at = datetime(2025, 6, 1)
        db.session.commit()

    # Only the January 2024 score falls in this range.
    resp = client.get(
        "/api/progress/me?start_date=2024-01-01&end_date=2024-01-31", headers=auth
    )
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["total_completed"] == 1
    assert data["total_points"] == 500

    # A range with no matching scores returns an empty summary.
    resp_empty = client.get(
        "/api/progress/me?start_date=2020-01-01&end_date=2020-01-31", headers=auth
    )
    assert resp_empty.get_json()["total_completed"] == 0

    # No range filter returns everything.
    resp_all = client.get("/api/progress/me", headers=auth)
    assert resp_all.get_json()["total_completed"] == 2


def test_progress_export_respects_date_range(client, student, app):
    from app.extensions import db
    from app.models import Score

    auth = auth_headers(student["token"])
    client.post("/api/scores", json={"game_id": 1, "points": 500, "completed": True}, headers=auth)

    with app.app_context():
        score = Score.query.first()
        score.created_at = datetime(2024, 1, 15)
        db.session.commit()

    import io

    from openpyxl import load_workbook

    resp_no_match = client.get(
        "/api/progress/export?start_date=2030-01-01&end_date=2030-01-31", headers=auth
    )
    assert resp_no_match.status_code == 200
    ws_no_match = load_workbook(io.BytesIO(resp_no_match.data)).active
    assert ws_no_match.max_row == 1  # yalnızca başlık satırı

    resp_with_match = client.get(
        "/api/progress/export?start_date=2024-01-01&end_date=2024-01-31", headers=auth
    )
    assert resp_with_match.status_code == 200
    ws_with_match = load_workbook(io.BytesIO(resp_with_match.data)).active
    assert ws_with_match.max_row == 2  # başlık + 1 skor satırı
    assert ws_with_match.cell(row=2, column=3).value == 500  # Puan sütunu
